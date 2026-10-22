#!/usr/bin/env python3
"""把原文变成 book.yaml 骨架，保证文案“搬运”而不是“重写”。
用法：python3 doc2yaml.py 原文.docx|.pdf|.md|.txt <project_dir>
输出：<project_dir>/doc.md（原文）和 <project_dir>/book.draft.yaml（每个标题一页，段落→p，列表→items）。
PPT / PDF 按 doc.md 里的页标记（<!-- 第 N 页 -->）把连续几页并成一个内容块（每块最多 PAGES_PER_BLOCK 页），免得 22 页原稿只产出 1 个块。
然后你只需：1) 给每页选 layout + image；2) 把相关段落挪进 on_photo/body/top/bottom；3) 加书写区 block。不要改写 p 里的文字。"""
import sys,os,re,subprocess,yaml
PAGES_PER_BLOCK=8   # 一个内容块最多合并几页 PPT（页标记切分）；6–8 页一块，lint 与版式分配消化得下
src,P=sys.argv[1],os.path.abspath(sys.argv[2]); os.makedirs(P,exist_ok=True)
def to_md(src):
    """.md/.txt 直接读；.docx/.pdf 等依次尝试 markitdown → pandoc → python-docx。"""
    if src.lower().endswith(('.md','.txt','.markdown')): return open(src,encoding='utf-8').read()
    for cmd in (['markitdown',src],['pandoc',src,'-t','gfm','--wrap=none']):
        try:
            r=subprocess.run(cmd,capture_output=True,text=True)
            if r.returncode==0 and r.stdout.strip(): return r.stdout
        except FileNotFoundError: pass
    if src.lower().endswith('.docx'):
        try:
            import docx; d=docx.Document(src); out=[]
            for p in d.paragraphs:
                st=(p.style.name or '').lower(); t=p.text.strip()
                if not t: out.append(''); continue
                lv=next((int(c) for c in st if c.isdigit()),1) if st.startswith(('heading','标题')) else 0
                out.append(('#'*min(lv,3)+' ' if lv else '- ' if 'list' in st else '')+t)
            for tb in d.tables:
                for row in tb.rows: out.append('| '+' | '.join(c.text.strip() for c in row.cells)+' |')
            return '\n'.join(out)
        except ImportError: pass
    sys.exit(f'无法转换 {src}：请安装 markitdown（pip install "markitdown[all]"）或 pandoc，或先把原文另存为 .md/.txt')
md=to_md(src)
open(os.path.join(P,'doc.md'),'w',encoding='utf-8').write(md)
pages=[]; cur=None; items=None; pend=None   # pend：extract.py 注释里记下的“原第 N 页 · 图：…”，挂到下一页的 note / image 上
pgno=0; blk0=0   # 当前原资料页号 / 当前内容块从第几页开始（按页标记控制每块页数，见 PAGES_PER_BLOCK）
def clean(s): return re.sub(r'\*\*|\\(?=[\\`*_{}\[\]()#+\-.!|>~<])','',s).strip()   # 去掉 markdown 粗体与转义（\# \. \- …）
for raw in md.splitlines():
    s=raw.strip()
    if not s: items=None; continue
    if re.fullmatch(r'!\[[^\]]*\]\([^)]*\)|<img\b[^>]*>',s): continue   # 文档里的图片（图片本身由 extract.py 抽到 photos/）
    cm=re.match(r'^<!--\s*(.*?)\s*-->$',s)
    if cm:
        pend=cm.group(1)
        pm=re.search(r'第\s*(\d+)\s*[页节]',pend)   # 页标记：<!-- 第 N 页 · 照片：… -->（Word 资料写“第 N 节”）
        if pm:
            k=int(pm.group(1))
            if cur is not None and k>pgno and k-blk0+1>PAGES_PER_BLOCK:   # 这个块已经到页数上限：就此收口，后面的页开新块
                cur=None; items=None; blk0=k
            elif cur is None: blk0=k
            pgno=k
        continue
    m=re.match(r'^(#{1,3})\s+(.*)',s)
    bold=re.fullmatch(r'\*\*(.+)\*\*',s)
    if m or (bold and len(bold.group(1))<(30 if re.search(r'[\u4e00-\u9fff]',bold.group(1)) else 70)):
        t=clean(m.group(2) if m else bold.group(1))
        im=re.search(r'照片：([^\s、]+)',pend or '')   # 只有照片才预填为版式图片；图形（截图、图表）要放 figure 块
        cur={'layout':'TODO','image':im.group(1) if im else 'TODO.jpg','header':['TODO EN','TODO'],'body':[{'title':t}]}
        if pend: cur={'note':'原'+pend,**cur}
        pages.append(cur); items=None; pend=None; blk0=pgno; continue
    if cur is None: cur={'layout':'POSTER','image':'TODO.jpg','top':[]}; pages.append(cur)
    key='body' if 'body' in cur else 'top'
    if s.startswith('|'):                       # markitdown 把 Word 文本框/表格变成 | … |
        cells=[c.strip() for c in s.strip('|').split('|')]
        if all(re.fullmatch(r'-*',c) for c in cells) or not any(cells): continue
        for c in cells:
            parts=[x.strip() for x in re.split(r'▶|•',c) if x.strip()]
            if not parts: continue
            hd=clean(parts[0]) if '▶' in c or '•' in c else None
            its=[clean(re.sub(r'(\\_)+|_{3,}','',x)) for x in (parts[1:] if hd else parts)]
            b={'items':[x for x in its if x]}
            if hd: b={'heading':hd,**b}
            cur[key].append(b)
        items=None; continue
    if re.fullmatch(r'(\\_|_)+',s): cur[key].append({'lines':1}); items=None; continue
    li=re.match(r'^([-*•▶✓]|\d+[.、)])\s*(.*)',s)
    if li:
        if items is None: items={'items':[]}; cur[key].append(items)
        items['items'].append(clean(li.group(2)))
    else:
        items=None; cur[key].append({'p':clean(s)})
class D(yaml.SafeDumper): pass
D.add_representer(str,lambda d,v:d.represent_scalar('tag:yaml.org,2002:str',v,style='"'))   # 生成的字符串统一双引号：撇号（Thailand's）不会炸解析
y=yaml.dump({'book':{'title':'TODO 手册标题','folio':'TODO 页脚文字','photos':'photos'},'pages':pages},Dumper=D,allow_unicode=True,sort_keys=False,width=1000)
y=('# 骨架：每个 TODO 都要替换。标准模式用内容版式 ONE/TWO/CARDS/STAGGER/FULL/COLLAGE/GRID/MIX（见 references/layouts.md，plan.md 给了每段的建议）；\n'
   '# 进阶模式用照片版式 PAPER/POSTER/T/B/SL/SR/HOLE/SPREAD_*（见 references/blocks.md）。\n'
   '# 相邻几个块可以合并到一页或一个跨页；文字只能挪动，不能改写。items 的 heading 是原文的小标题。\n'+y)
open(os.path.join(P,'book.draft.yaml'),'w',encoding='utf-8').write(y)
print(f'{len(pages)} 个内容块 → {P}/book.draft.yaml ；原文 → {P}/doc.md')
